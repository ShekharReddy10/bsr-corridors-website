"""Excel exports and JSON backups."""

import io
from datetime import date

from django.core import serializers
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import AuditLog, Guest, Room, RoomType, Stay, StayGuest

HEADER_FILL = PatternFill("solid", fgColor="1F2A30")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    return ws


def _id(value: str, full: bool) -> str:
    if not value:
        return ""
    return value if full else f"••••{value[-4:]}"


def stays_in_range(start: date | None = None, end: date | None = None):
    stays = (Stay.objects.select_related("guest", "room", "room__room_type").prefetch_related("other_guests")
             .order_by("check_in", "room__number"))
    if start and end:
        stays = stays.filter(check_in__lte=end, check_out__gte=start)
    return stays


def build_workbook(start: date | None = None, end: date | None = None, full_ids: bool = False) -> bytes:
    stays = list(stays_in_range(start, end))
    wb = Workbook()
    _sheet(
        wb,
        "Stays",
        ["Stay ID", "Guest", "Phone", "Room", "Room type", "Check-in", "Check-out", "Nights", "Guests",
         "Booked via", "Booking ref", "Total", "Paid", "Balance", "Payment mode", "Status", "Nationality",
         "ID proof", "ID number", "Address", "Other guests", "Form C filed", "Notes"],
        [
            [s.pk, s.guest.name, s.guest.phone, s.room.number, s.room.room_type.name, s.check_in, s.check_out,
             s.nights, s.num_guests, s.get_source_display(), s.source_ref, float(s.total), float(s.amount_paid), float(s.balance),
             s.get_payment_mode_display(), s.get_status_display(), s.guest.nationality,
             s.guest.get_id_type_display(), _id(s.guest.id_number, full_ids), s.guest.address,
             "; ".join(f"{o.name or 'Guest ' + str(o.position)}"
                       + (f", {o.phone}" if o.phone else "")
                       + (f" ({o.get_id_type_display()} {_id(o.id_number, full_ids)})" if o.id_number else "")
                       for o in s.other_guests.all()),
             "Yes" if s.form_c_filed else ("No" if s.guest.is_foreign else ""), s.notes]
            for s in stays
        ],
        first=True,
    )
    guests = {s.guest_id: s.guest for s in stays} if start else {g.pk: g for g in Guest.objects.all()}
    _sheet(
        wb,
        "Guests",
        ["Guest ID", "Name", "Phone", "Address", "Nationality", "ID proof", "ID number",
         "Passport", "Visa number", "Visa type", "Arrival in India"],
        [
            [g.pk, g.name, g.phone, g.address, g.nationality, g.get_id_type_display(), _id(g.id_number, full_ids),
             _id(g.passport_number, full_ids), g.visa_number, g.visa_type, g.arrival_in_india]
            for g in sorted(guests.values(), key=lambda g: g.name.lower())
        ],
    )
    _sheet(
        wb,
        "Rooms",
        ["Room", "Floor", "Type", "AC", "Bed", "Max guests", "Status", "Active"],
        [
            [r.number, r.floor, r.room_type.name, "AC" if r.room_type.is_ac else "Non-AC", r.room_type.bed_type,
             r.room_type.max_guests, r.get_status_display(), "Yes" if r.is_active else "No"]
            for r in Room.objects.select_related("room_type")
        ],
    )
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "DD-MM-YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BACKUP_MODELS = [RoomType, Room, Guest, Stay, StayGuest, AuditLog]


def backup_json() -> bytes:
    """Full backup in Django fixture format (restore with `manage.py loaddata`).

    ID numbers stay encrypted — restoring needs the same FIELD_ENCRYPTION_KEY.
    """
    objects = [obj for model in BACKUP_MODELS for obj in model.objects.order_by("pk")]
    return serializers.serialize("json", objects, indent=1).encode()
