import json
from datetime import timedelta
from decimal import Decimal
from unittest import skipUnless

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import IntegrityError, connection, transaction
from django.test import TestCase, override_settings

from core import services
from core.exports import backup_json, build_workbook
from core.models import AuditLog, Guest, Room, RoomType, Stay, StayGuest

T = services.today()


def make_rooms():
    ac = RoomType.objects.create(name="Luxury AC", is_ac=True, max_guests=2)
    non = RoomType.objects.create(name="Luxury Non-AC", is_ac=False, max_guests=2)
    rooms = {
        "101": Room.objects.create(number="101", room_type=ac, sort_order=1),
        "102": Room.objects.create(number="102", room_type=ac, sort_order=2),
        "103": Room.objects.create(number="103", room_type=ac, sort_order=3),
        "201": Room.objects.create(number="201", room_type=non, sort_order=4),
    }
    return ac, non, rooms


def make_guest(name="Ravi Kumar", phone="+91 98765 43210", aadhaar="123412341234"):
    return Guest.objects.create(name=name, phone=phone, address="Gachibowli, Hyderabad", id_number=aadhaar)


def stay(guest, room, start_offset, nights, **kw):
    s = Stay(guest=guest, room=room, check_in=T + timedelta(days=start_offset),
             check_out=T + timedelta(days=start_offset + nights), total_amount=Decimal(1500 * nights), **kw)
    return services.save_stay(s, action="create", summary="test")


class EncryptionTests(TestCase):
    def test_id_number_encrypted_at_rest_and_in_backup(self):
        g = make_guest()
        with connection.cursor() as cur:
            cur.execute(f"SELECT id_number FROM {Guest._meta.db_table} WHERE id = %s", [g.pk])
            raw = cur.fetchone()[0]
        self.assertTrue(raw.startswith("enc:"))
        self.assertNotIn("123412341234", raw)
        self.assertEqual(Guest.objects.get(pk=g.pk).id_number, "123412341234")
        self.assertEqual(g.id_masked, "•••• •••• 1234")
        dump = backup_json().decode()
        self.assertNotIn("123412341234", dump)
        self.assertIn("enc:", dump)

    def test_backup_round_trip(self):
        _, _, rooms = make_rooms()
        g = make_guest()
        stay(g, rooms["101"], 0, 2)
        data = backup_json()
        Stay.objects.all().delete()
        Guest.objects.all().delete()
        from django.core import serializers
        for obj in serializers.deserialize("json", data):
            obj.save()
        self.assertEqual(Guest.objects.get().id_number, "123412341234")
        self.assertEqual(Stay.objects.count(), 1)


class AvailabilityTests(TestCase):
    def setUp(self):
        _, _, self.rooms = make_rooms()
        self.g = make_guest()

    def test_overlap_rejected_back_to_back_allowed(self):
        stay(self.g, self.rooms["101"], 0, 3)
        with self.assertRaises(services.StayConflict):
            stay(make_guest("Anita", "+91 90000 00001"), self.rooms["101"], 2, 2)
        stay(make_guest("Anita", "+91 90000 00011"), self.rooms["101"], 3, 2)  # arrives on checkout day
        self.assertEqual(Stay.objects.count(), 2)

    def test_cancelled_stay_frees_room(self):
        s = stay(self.g, self.rooms["101"], 0, 3)
        services.cancel(s)
        stay(make_guest("Anita", "+91 90000 00001"), self.rooms["101"], 1, 1)

    def test_transfer_on_arrival_frees_room(self):
        s = stay(self.g, self.rooms["201"], 0, 3)
        services.transfer_out(s, "  Hotel Sitara  ", "Wanted AC room")
        s.refresh_from_db()
        self.assertEqual(s.status, Stay.Status.TRANSFERRED)
        self.assertEqual((s.transferred_to, s.transfer_reason), ("Hotel Sitara", "Wanted AC room"))
        self.assertIn("201", services.free_rooms(T, T + timedelta(days=3)).values_list("number", flat=True))
        stay(make_guest("Anita", "+91 90000 00001"), self.rooms["201"], 0, 3)  # DB constraint lets it through too
        self.assertIn("Hotel Sitara", AuditLog.objects.filter(action="transfer").get().summary)

    def test_transfer_after_nights_stayed_keeps_them(self):
        s = stay(self.g, self.rooms["201"], -2, 5, status=Stay.Status.CHECKED_IN)
        services.transfer_out(s, "Hotel Sitara")
        s.refresh_from_db()
        self.assertEqual((s.status, s.check_out, s.transferred_to), (Stay.Status.CHECKED_OUT, T, "Hotel Sitara"))
        stay(make_guest("Anita", "+91 90000 00001"), self.rooms["201"], 0, 2)
        with self.assertRaises(services.StayConflict):
            stay(make_guest("Vijay", "+91 90000 00002"), self.rooms["201"], -1, 1)

    def test_undo_early_check_out_restores_booked_date(self):
        s = stay(self.g, self.rooms["201"], -1, 4, status=Stay.Status.CHECKED_IN)
        booked = s.check_out
        services.check_out(s)
        self.assertEqual((s.check_out, services.booked_check_out(s)), (T, booked))
        services.undo_check_out(s, services.booked_check_out(s))
        s.refresh_from_db()
        self.assertEqual((s.status, s.check_out, s.checked_out_at), (Stay.Status.CHECKED_IN, booked, None))

    def test_undo_check_out_blocked_if_room_rebooked(self):
        s = stay(self.g, self.rooms["201"], -1, 4, status=Stay.Status.CHECKED_IN)
        booked = s.check_out
        services.check_out(s)
        stay(make_guest("Anita", "+91 90000 00001"), self.rooms["201"], 1, 1)
        with self.assertRaises(services.StayConflict):
            services.undo_check_out(s, booked)

    def test_undo_monthly_check_out_reopens(self):
        s = Stay(guest=self.g, room=self.rooms["201"], kind=Stay.Kind.MONTHLY, check_in=T - timedelta(days=5),
                 status=Stay.Status.CHECKED_IN)
        services.save_stay(s, action="create", summary="test")
        services.check_out(s)
        self.assertIsNone(services.booked_check_out(s))
        services.undo_check_out(s, None)
        s.refresh_from_db()
        self.assertEqual((s.status, s.check_out), (Stay.Status.CHECKED_IN, None))

    def test_undo_check_in(self):
        s = stay(self.g, self.rooms["201"], 0, 2)
        services.check_in(s)
        services.undo_check_in(s)
        s.refresh_from_db()
        self.assertEqual((s.status, s.checked_in_at), (Stay.Status.UPCOMING, None))

    def test_transfer_needs_hotel_name(self):
        s = stay(self.g, self.rooms["201"], 0, 1)
        with self.assertRaises(ValueError):
            services.transfer_out(s, "  ")
        s.refresh_from_db()
        self.assertEqual(s.status, Stay.Status.UPCOMING)

    @skipUnless(connection.vendor == "postgresql", "exclusion constraint is Postgres-only")
    def test_database_blocks_overlap_even_without_service(self):
        stay(self.g, self.rooms["101"], 0, 3)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Stay.objects.create(guest=self.g, room=self.rooms["101"], check_in=T + timedelta(days=1),
                                check_out=T + timedelta(days=2))

    def test_free_rooms(self):
        stay(self.g, self.rooms["101"], 0, 3)
        free = set(services.free_rooms(T, T + timedelta(days=1)).values_list("number", flat=True))
        self.assertEqual(free, {"102", "103", "201"})


class GuestPhoneTests(TestCase):
    def test_one_guest_per_number(self):
        make_guest(phone="+91 98765 43210")
        with self.assertRaises(IntegrityError), transaction.atomic():
            make_guest("Other", phone="098765 43210")  # same last 10 digits
        self.assertEqual(Guest.objects.get().phone_key, "9876543210")


class ExtensionTests(TestCase):
    def setUp(self):
        self.ac, _, self.rooms = make_rooms()
        self.g = make_guest()

    def test_extend_in_place_when_free(self):
        s = stay(self.g, self.rooms["101"], 0, 2)
        plan = services.plan_extension(s, s.check_out + timedelta(days=2))
        self.assertTrue(plan.free)
        services.extend_in_place(s, plan.new_check_out, Decimal("2500"), Decimal("1000"), "upi")
        s.refresh_from_db()
        self.assertEqual(s.nights, 4)
        self.assertEqual((s.total, s.amount_paid, s.balance), (Decimal("5500"), Decimal("1000"), Decimal("4500")))
        self.assertTrue(AuditLog.objects.filter(stay=s, action="extend").exists())

    def test_type_has_space_moves_future_booking(self):
        """Room 101 is booked after; another AC room is free → extend in 101 and move the future booking."""
        s = stay(self.g, self.rooms["101"], 0, 2)
        nxt = stay(make_guest("Next", "+91 90000 00002"), self.rooms["101"], 2, 3)
        stay(make_guest("Busy", "+91 90000 00003"), self.rooms["102"], 2, 1)  # 102 busy on first extra night
        plan = services.plan_extension(s, s.check_out + timedelta(days=2))
        self.assertFalse(plan.free)
        self.assertEqual(plan.full_nights, [])
        self.assertTrue(plan.can_extend_in_place)
        self.assertEqual([(b.pk, r.number) for b, r in plan.moves], [(nxt.pk, "103")])  # 103 free for Next's whole stay
        services.extend_with_moves(s, plan.new_check_out, plan.moves, Decimal("3000"))
        s.refresh_from_db(); nxt.refresh_from_db()
        self.assertEqual((s.room.number, s.nights, s.total), ("101", 4, Decimal("6000")))
        self.assertEqual(nxt.room.number, "103")
        self.assertTrue(AuditLog.objects.filter(stay=nxt, action="move").exists())

    def test_type_full_offers_options(self):
        """All AC rooms taken on an extra night → no automatic extension; split/move options offered."""
        s = stay(self.g, self.rooms["101"], 0, 2)
        stay(make_guest("Next", "+91 90000 00002"), self.rooms["101"], 2, 3)
        stay(make_guest("B2", "+91 90000 00003"), self.rooms["102"], 2, 1)
        stay(make_guest("B3", "+91 90000 00004"), self.rooms["103"], 2, 1)
        plan = services.plan_extension(s, s.check_out + timedelta(days=2))
        self.assertEqual(plan.full_nights, [T + timedelta(days=2)])
        self.assertFalse(plan.can_extend_in_place)
        self.assertEqual(plan.split_rooms, [])  # no AC room free for the extra nights
        self.assertEqual([r.number for r in plan.move_rooms], ["201"])  # other type offered
        cont = services.extend_split(s, plan.new_check_out, self.rooms["201"], Decimal("2000"))
        self.assertEqual((cont.check_in, cont.total, cont.linked_to), (s.check_out, Decimal("2000"), s))
        self.assertEqual(len(s.chain()), 2)

    def test_in_house_guest_never_moved_automatically(self):
        s = stay(self.g, self.rooms["101"], 0, 2)
        other = stay(make_guest("In", "+91 90000 00005"), self.rooms["102"], 0, 5)
        services.check_in(other)
        # nothing blocks 101 itself → simple extension
        self.assertTrue(services.plan_extension(s, s.check_out + timedelta(days=1)).free)

    def test_move_mid_stay_splits_at_today(self):
        s = stay(self.g, self.rooms["101"], -2, 3)  # arrived 2 days ago, leaves tomorrow
        services.check_in(s)
        stay(make_guest("Next", "+91 90000 00002"), self.rooms["101"], 1, 3)
        cont = services.extend_move(s, T + timedelta(days=4), self.rooms["102"], T)
        s.refresh_from_db()
        self.assertEqual(s.check_out, T)
        self.assertEqual(s.status, Stay.Status.CHECKED_OUT)
        self.assertEqual((cont.room.number, cont.check_in, cont.status), ("102", T, Stay.Status.CHECKED_IN))

    def test_early_checkout_shortens(self):
        s = stay(self.g, self.rooms["101"], -1, 4)
        services.check_in(s)
        services.check_out(s)
        s.refresh_from_db()
        self.assertEqual((s.check_out, s.status), (T, Stay.Status.CHECKED_OUT))

    def test_payment(self):
        s = stay(self.g, self.rooms["101"], 0, 2)
        services.add_payment(s, Decimal("1000"), "upi")
        s.refresh_from_db()
        self.assertEqual((s.amount_paid, s.balance, s.payment_mode), (Decimal("1000"), Decimal("2000"), "upi"))


class ExportTests(TestCase):
    def test_workbook_masks_ids_by_default(self):
        _, _, rooms = make_rooms()
        stay(make_guest(), rooms["101"], 0, 1)
        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(build_workbook(T, T + timedelta(days=5))))
        values = [c.value for row in wb["Stays"].iter_rows() for c in row]
        self.assertIn("••••1234", values)
        self.assertNotIn("123412341234", values)
        self.assertEqual(wb.sheetnames, ["Stays", "Guests", "Rooms"])


class ViewTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = get_user_model().objects.create_user("test-admin", password="test-pass-123456")
        self.client.force_login(self.user)
        self.ac, _, self.rooms = make_rooms()

    def _new_stay_post(self, **over):
        return self.client.post("/stays/new/", {**self._post_data(edit=False), **over})

    def _post_data(self, edit=True):
        stay = Stay.objects.first() if edit else None
        data = {
            "g-phone": "+91 98765 43210", "g-name": "Ravi Kumar", "g-address": "Hyderabad", "g-nationality": "Indian",
            "g-id_type": "aadhaar", "g-id_number": "1234 1234 1234",
            "s-room": self.rooms["101"].pk, "s-check_in": T.isoformat(), "s-check_out": (T + timedelta(days=2)).isoformat(),
            "s-num_guests": 2, "s-total_amount": "3000", "s-amount_paid": "1000", "s-payment_mode": "upi",
            "s-source": "airbnb", "s-source_ref": "HMABC123",
            "o-TOTAL_FORMS": "2", "o-INITIAL_FORMS": "0", "o-MIN_NUM_FORMS": "0", "o-MAX_NUM_FORMS": "9",
            "o-0-id_type": "aadhaar", "o-1-id_type": "aadhaar",
        }
        if stay:  # editing: keep the same guest record
            data["g-guest_id"] = stay.guest_id
        return data

    def test_pages_render(self):
        s = stay(make_guest(), self.rooms["101"], 0, 2)
        for url in ["/", "/calendar/", "/calendar/?days=31", f"/calendar/room/{self.rooms['101'].pk}/", "/stays/",
                    "/stays/new/", f"/stays/{s.pk}/", f"/stays/{s.pk}/edit/", f"/stays/{s.pk}/extend/", "/rooms/",
                    "/rooms/type/new/", "/rooms/room/new/", f"/rooms/room/{self.rooms['101'].pk}/", "/export/",
                    "/backups/", "/settings/"]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
        r = self.client.get("/stays/?q=ravi", HTTP_HX_REQUEST="true")
        self.assertContains(r, "Ravi Kumar")
        self.assertNotContains(r, "<html")

    def test_create_stay_and_returning_guest(self):
        r = self._new_stay_post()
        s = Stay.objects.get()
        self.assertRedirects(r, f"/stays/{s.pk}/")
        self.assertEqual(s.guest.id_number, "123412341234")
        self.assertEqual(s.balance, Decimal("2000"))
        self.assertEqual((s.source, s.source_ref), ("airbnb", "HMABC123"))

        # Same number without picking the saved guest → asked to use the saved guest
        r = self._new_stay_post(**{"g-name": "Someone Else", "s-room": self.rooms["103"].pk})
        self.assertContains(r, "already saved for Ravi Kumar")

        lookup = self.client.get("/guests/lookup/?phone=9876543210").json()
        self.assertTrue(lookup["found"])
        self.assertNotIn("123412341234", json.dumps(lookup))

        # Returning guest, ID left blank → stored ID kept, same guest record
        r = self._new_stay_post(**{"g-guest_id": lookup["guest_id"], "g-id_number": "", "s-room": self.rooms["102"].pk})
        self.assertEqual(Guest.objects.count(), 1)
        self.assertEqual(Guest.objects.get().id_number, "123412341234")

    def test_conflict_shows_error(self):
        self._new_stay_post()
        r = self._new_stay_post(**{"g-phone": "+91 90000 00009", "g-name": "Other"})
        self.assertContains(r, "already booked")
        self.assertEqual(Stay.objects.count(), 1)

    def test_validation(self):
        r = self._new_stay_post(**{"g-id_number": "1234"})
        self.assertContains(r, "Aadhaar number must be 12 digits")
        r = self._new_stay_post(**{"g-id_number": "", "g-phone": "+91 90000 00077"})  # ID number optional
        self.assertEqual(Stay.objects.get().guest.id_number, "")
        self.assertContains(self.client.get(f"/stays/{Stay.objects.get().pk}/"), "Not provided")
        Stay.objects.all().delete(); Guest.objects.all().delete()
        # Foreign national details are required at check-in (not when booking ahead)
        r = self._new_stay_post(**{"g-nationality": "British", "g-id_type": "passport", "g-id_number": "X1",
                                   "g-phone": "+44 7700 900111", "checkin": "1"})
        self.assertContains(r, "Required for foreign nationals")
        r = self._new_stay_post(**{"s-num_guests": 3})
        self.assertContains(r, "allows up to 2 guests")
        self.assertEqual(Stay.objects.count(), 0)

    def test_extend_flow(self):
        s = stay(make_guest(), self.rooms["101"], 0, 2)
        r = self.client.post(f"/stays/{s.pk}/extend/", {"new_check_out": (s.check_out + timedelta(days=1)).isoformat(),
                                                         "amount": "1200", "paid_now": "1200", "payment_mode": "cash"})
        self.assertRedirects(r, f"/stays/{s.pk}/")
        s.refresh_from_db()
        self.assertEqual((s.total, s.amount_paid), (Decimal("4200"), Decimal("1200")))
        # Room 101 booked by a future guest, but other AC rooms are free → extended, future booking moved
        nxt = stay(make_guest("Next", "+91 90000 00002"), self.rooms["101"], 3, 2)
        r = self.client.post(f"/stays/{s.pk}/extend/", {"new_check_out": (T + timedelta(days=5)).isoformat(), "amount": "2000"},
                             follow=True)
        self.assertContains(r, "Next’s booking moved from Room 101 to Room 102")
        s.refresh_from_db(); nxt.refresh_from_db()
        self.assertEqual((s.room.number, s.check_out, nxt.room.number), ("101", T + timedelta(days=5), "102"))

        # Every AC room taken on the next extra night → options page instead
        stay(make_guest("Block1", "+91 90000 00006"), self.rooms["101"], 5, 2)
        stay(make_guest("Block2", "+91 90000 00007"), self.rooms["102"], 5, 2)
        stay(make_guest("Block3", "+91 90000 00008"), self.rooms["103"], 5, 2)
        r = self.client.post(f"/stays/{s.pk}/extend/", {"new_check_out": (T + timedelta(days=6)).isoformat(), "amount": "1500"})
        self.assertContains(r, "All Luxury AC rooms are booked")
        self.assertContains(r, 'name="amount" value="1500')
        self.assertContains(r, "Room <b>201</b>")
        r = self.client.post(f"/stays/{s.pk}/extend/", {"new_check_out": (T + timedelta(days=6)).isoformat(), "amount": "1500",
                                                         "choice": "split", "room": self.rooms["201"].pk})
        cont = Stay.objects.get(linked_to=s)
        self.assertEqual((cont.room.number, cont.total, cont.source), ("201", Decimal("1500"), s.source))

    def test_actions_and_reveal(self):
        s = stay(make_guest(), self.rooms["101"], 0, 2)
        self.client.post(f"/stays/{s.pk}/check-in/")
        self.client.post(f"/stays/{s.pk}/payment/", {"amount": "500", "mode": "cash"})
        s.refresh_from_db()
        self.assertEqual((s.status, s.amount_paid), (Stay.Status.CHECKED_IN, Decimal("500")))
        r = self.client.post(f"/stays/{s.pk}/reveal-id/", {"field": "id_number"})
        self.assertContains(r, "123412341234")
        self.assertTrue(AuditLog.objects.filter(action="reveal_id").exists())

    def test_other_guests_ids_optional(self):
        RoomType.objects.filter(pk=self.ac.pk).update(max_guests=3)
        # 2 guests: second guest's ID saved (encrypted), third row ignored
        r = self._new_stay_post(**{"o-0-name": "Sita Kumar", "o-0-phone": "+91 91234 56789", "o-0-id_number": "5678 5678 5678",
                                   "o-1-name": "Ignored", "o-1-id_number": "1111"})
        s = Stay.objects.get()
        others = list(s.other_guests.all())
        self.assertEqual([(o.position, o.name, o.phone, o.id_number) for o in others],
                         [(2, "Sita Kumar", "+91 91234 56789", "567856785678")])
        with connection.cursor() as cur:
            cur.execute(f"SELECT id_number FROM {StayGuest._meta.db_table}")
            self.assertTrue(cur.fetchone()[0].startswith("enc:"))
        page = self.client.get(f"/stays/{s.pk}/")
        self.assertContains(page, "Sita Kumar")
        self.assertContains(page, 'href="tel:+91 91234 56789"')
        self.assertNotContains(page, "567856785678")
        r = self.client.post(f"/stays/{s.pk}/reveal-id/", {"field": f"other:{others[0].pk}"})
        self.assertContains(r, "567856785678")

        # Edit: 3 guests, guest 2 ID left blank → kept; guest 3 name only (ID optional)
        o = others[0]
        r = self.client.post(f"/stays/{s.pk}/edit/", {
            **{k: v for k, v in self._post_data().items()}, "s-num_guests": 3,
            "o-TOTAL_FORMS": "2", "o-INITIAL_FORMS": "1", "o-0-id": o.pk, "o-0-stay": s.pk, "o-0-name": "Sita Kumar",
            "o-0-phone": "+91 91234 56789", "o-0-id_type": "aadhaar", "o-0-id_number": "",
            "o-1-name": "", "o-1-phone": "+91 99999 00000", "o-1-id_type": "aadhaar", "o-1-id_number": "",
        })
        self.assertRedirects(r, f"/stays/{s.pk}/")
        self.assertEqual([(x.position, x.name, x.id_number) for x in s.other_guests.all()],
                         [(2, "Sita Kumar", "567856785678"), (3, "", "")])
        self.assertEqual(s.other_guests.get(position=3).phone, "+91 99999 00000")  # phone only is enough

        # Bad Aadhaar on an extra guest is rejected
        r = self.client.post(f"/stays/{s.pk}/edit/", {**self._post_data(), "s-num_guests": 2, "o-TOTAL_FORMS": "1",
                                                      "o-INITIAL_FORMS": "0", "o-0-name": "X", "o-0-id_type": "aadhaar",
                                                      "o-0-id_number": "12"})
        self.assertContains(r, "Aadhaar number must be 12 digits")
        r = self.client.post(f"/stays/{s.pk}/edit/", {**self._post_data(), "s-num_guests": 2, "o-TOTAL_FORMS": "1",
                                                      "o-INITIAL_FORMS": "0", "o-0-phone": "123"})
        self.assertContains(r, "at least 10 digits")

        # Back to 1 guest → extra guests removed
        r = self.client.post(f"/stays/{s.pk}/edit/", {**self._post_data(), "s-num_guests": 1})
        self.assertFalse(StayGuest.objects.exists())

    def test_dashboard_tomorrow(self):
        leaving = stay(make_guest(), self.rooms["101"], -1, 2)           # leaves tomorrow → 101 free tomorrow night
        stay(make_guest("Anita", "+91 90000 00021"), self.rooms["102"], 0, 3)   # stays through tomorrow
        stay(make_guest("Kiran", "+91 90000 00022"), self.rooms["103"], 1, 2)   # arrives tomorrow
        r = self.client.get("/")
        self.assertEqual([s.pk for s in r.context["departing_tomorrow"]], [leaving.pk])
        self.assertEqual(sorted(x.number for x in r.context["vacant_tomorrow"]), ["101", "201"])
        self.assertContains(r, "Vacant tomorrow night")

    def _quick(self, name, room, start, nights=None, **over):
        data = {"g-name": name, "g-nationality": "Indian", "g-id_type": "aadhaar",
                "s-source": "airbnb", "s-kind": "daily", "s-room": room.pk, "s-check_in": (T + timedelta(days=start)).isoformat(),
                "s-check_out": (T + timedelta(days=start + nights)).isoformat() if nights else "",
                "s-num_guests": 1, "s-total_amount": "2000", "s-amount_paid": "500", "o-TOTAL_FORMS": "0", "o-INITIAL_FORMS": "0"}
        data.update(over)
        return self.client.post("/stays/new/", data)

    def test_quick_future_booking_then_check_in(self):
        self._quick("Meena", self.rooms["101"], 5, 2)
        self._quick("Arun", self.rooms["102"], 1, 3)          # second guest without phone: allowed
        meena = Stay.objects.get(guest__name="Meena")
        self.assertEqual((meena.guest.phone, meena.guest.phone_key, meena.amount_paid), ("", None, Decimal("500")))
        # Check in without details → asked to add them
        r = self.client.post(f"/stays/{meena.pk}/check-in/")
        self.assertRedirects(r, f"/stays/{meena.pk}/edit/?checkin=1", fetch_redirect_response=False)
        page = self.client.get(f"/stays/{meena.pk}/edit/?checkin=1")
        self.assertContains(page, "Save &amp; check in")
        base = {**self._quick_data(meena), "checkin": "1"}
        r = self.client.post(f"/stays/{meena.pk}/edit/", base)
        self.assertContains(r, "This field is required")
        r = self.client.post(f"/stays/{meena.pk}/edit/", {**base, "g-phone": "+91 90000 12345", "g-address": "Kukatpally"})
        meena.refresh_from_db()
        self.assertEqual(meena.status, Stay.Status.CHECKED_IN)

    def _quick_data(self, stay):
        return {"g-guest_id": stay.guest_id, "g-name": stay.guest.name, "g-nationality": "Indian", "g-id_type": "aadhaar",
                "s-source": stay.source, "s-kind": stay.kind, "s-room": stay.room_id, "s-check_in": stay.check_in.isoformat(),
                "s-check_out": stay.check_out.isoformat() if stay.check_out else "", "s-num_guests": 1,
                "s-total_amount": str(stay.total_amount), "s-amount_paid": str(stay.amount_paid),
                "o-TOTAL_FORMS": "0", "o-INITIAL_FORMS": "0"}

    def test_monthly_open_ended(self):
        r = self._quick("Long Stay", self.rooms["201"], -10, **{"s-kind": "monthly", "s-monthly_rent": "12000",
                                                               "s-deposit_amount": "5000", "s-total_amount": "12000"})
        m = Stay.objects.get()
        self.assertIsNone(m.check_out)
        self.assertEqual((m.nights, m.deposit_amount), (10, Decimal("5000")))
        # Daily stay needs a check-out
        r = self._quick("No Date", self.rooms["103"], 1)
        self.assertContains(r, "choose “Monthly”")
        # Room blocked far in the future while open-ended
        r = self._quick("Later", self.rooms["201"], 60, 2)
        self.assertContains(r, "already booked")
        self.assertIn(self.rooms["201"].pk, [b.room_id for b in Stay.objects.live().overlapping(T + timedelta(days=400))])
        # Pages render with an open-ended stay
        for url in ["/", "/calendar/", f"/calendar/room/{self.rooms['201'].pk}/", "/stays/", f"/stays/{m.pk}/", "/upcoming/"]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(f"/stays/{m.pk}/check-in/")  # no phone → redirected, still upcoming
        m.guest.phone, m.guest.address = "+91 90000 33333", "Hyd"; m.guest.save()
        self.client.post(f"/stays/{m.pk}/check-in/")
        self.client.post(f"/stays/{m.pk}/add-rent/")
        m.refresh_from_db()
        self.assertEqual((m.status, m.total), (Stay.Status.CHECKED_IN, Decimal("24000")))
        r = self.client.post("/export/", {"start": T.isoformat(), "end": T.isoformat(), "action": "download"})
        self.assertEqual(r.status_code, 200)
        self.client.post(f"/stays/{m.pk}/check-out/")
        m.refresh_from_db()
        self.assertEqual((m.check_out, m.status), (T, Stay.Status.CHECKED_OUT))

    def test_upcoming_lists(self):
        a = self._quick("Tomorrow Guest", self.rooms["101"], 1, 2)
        self._quick("Day After", self.rooms["102"], 2, 1)
        self._quick("Next Month", self.rooms["103"], 30, 3)
        r = self.client.get("/")
        self.assertEqual([s.guest.name for s in r.context["arriving_tomorrow"]], ["Tomorrow Guest"])
        self.assertEqual([s.guest.name for s in r.context["arriving_day_after"]], ["Day After"])
        self.assertEqual(r.context["upcoming_count"], 3)
        r = self.client.get("/upcoming/")
        for name in ["Tomorrow Guest", "Day After", "Next Month"]:
            self.assertContains(r, name)

    def test_guest_search_and_rebook(self):
        g = make_guest()
        stay(g, self.rooms["101"], -10, 2)
        r = self.client.get("/guests/?q=98765", HTTP_HX_REQUEST="true")
        self.assertContains(r, "Ravi Kumar")
        self.assertContains(r, f"/stays/new/?guest={g.pk}")
        self.assertContains(self.client.get(f"/guests/{g.pk}/"), "New booking for Ravi Kumar")
        r = self.client.get(f"/stays/new/?guest={g.pk}")
        self.assertContains(r, "Booking for <b>Ravi Kumar</b>")
        self.assertContains(r, f'name="g-guest_id" value="{g.pk}"')
        self.assertEqual(self.client.get("/guests/").status_code, 200)

    def test_room_with_history_is_deactivated_not_deleted(self):
        stay(make_guest(), self.rooms["101"], 0, 1)
        self.client.post(f"/rooms/room/{self.rooms['101'].pk}/delete/")
        self.assertFalse(Room.objects.get(pk=self.rooms["101"].pk).is_active)
        self.client.post(f"/rooms/room/{self.rooms['103'].pk}/delete/")
        self.assertFalse(Room.objects.filter(pk=self.rooms["103"].pk).exists())

    def test_export_download(self):
        stay(make_guest(), self.rooms["101"], 0, 1)
        r = self.client.post("/export/", {"start": T.isoformat(), "end": T.isoformat(), "action": "download"})
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_purge(self):
        old = stay(make_guest(), self.rooms["101"], -2000, 2)
        old.status = Stay.Status.CHECKED_OUT
        old.save()
        self.client.post("/settings/", {"years": "3", "confirm": "PURGE"})
        self.assertFalse(Stay.objects.exists())
        self.assertFalse(Guest.objects.exists())


class SecurityTests(TestCase):
    def setUp(self):
        cache.clear()
        get_user_model().objects.create_user("test-admin", password="test-pass-123456")

    def test_login_required(self):
        self.assertRedirects(self.client.get("/stays/"), "/login/?next=/stays/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/healthz/").status_code, 200)

    def test_login_and_lockout(self):
        r = self.client.post("/login/", {"username": "test-admin", "password": "test-pass-123456"})
        self.assertRedirects(r, "/", fetch_redirect_response=False)
        self.client.logout()
        for _ in range(5):
            self.assertContains(self.client.post("/login/", {"username": "test-admin", "password": "wrong"}),
                                "Incorrect username or password")
        r = self.client.post("/login/", {"username": "test-admin", "password": "test-pass-123456"})
        self.assertContains(r, "Too many failed attempts")
        self.assertNotIn("_auth_user_id", self.client.session)

    @override_settings(BACKUP_TOKEN="secret-token")
    def test_backup_hook_requires_token(self):
        self.assertEqual(self.client.post("/backup/run/").status_code, 403)
        self.assertEqual(self.client.post("/backup/run/", HTTP_X_BACKUP_TOKEN="nope").status_code, 403)
        r = self.client.post("/backup/run/", HTTP_X_BACKUP_TOKEN="secret-token")  # Drive not configured → 500 + message
        self.assertEqual(r.status_code, 500)
        self.assertIn("not connected", r.json()["message"])
